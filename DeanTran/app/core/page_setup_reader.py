"""
Page Setup Reader - Reads margins, paper size, and orientation from Word/Excel/PPT documents.
"""
from __future__ import annotations

import logging
from pathlib import Path
from app.core.pdf_exporter import PAPER_SIZES, _OPENPYXL_PAPER, PageSetup

logger = logging.getLogger("DeanTran.page_setup_reader")


def to_cm(val) -> float:
    """Convert python-docx length to centimetres. 1 cm = 360000 EMU."""
    if val is None:
        return 0.0
    if hasattr(val, "cm"):
        return round(val.cm, 2)
    return round(float(val) / 360000.0, 2)


def inch_to_cm(val) -> float:
    """Convert inches to centimetres."""
    if val is None:
        return 0.0
    return round(float(val) * 2.54, 2)


def read_page_setup_docx(docx_path: Path) -> PageSetup:
    """Read page setup from a Word document (.docx)."""
    from docx import Document
    from docx.enum.section import WD_ORIENT

    doc = Document(str(docx_path))
    if not doc.sections:
        return PageSetup()

    section = doc.sections[0]
    w_cm = to_cm(section.page_width)
    h_cm = to_cm(section.page_height)

    orient = "portrait"
    if section.orientation == WD_ORIENT.LANDSCAPE:
        orient = "landscape"
    elif section.orientation == WD_ORIENT.PORTRAIT:
        orient = "portrait"
    else:
        if w_cm > h_cm:
            orient = "landscape"

    # Match paper size
    port_w, port_h = (w_cm, h_cm) if orient == "portrait" else (h_cm, w_cm)
    matched_paper = "Custom"
    for name, (pw, ph) in PAPER_SIZES.items():
        if abs(pw - port_w) < 0.25 and abs(ph - port_h) < 0.25:
            matched_paper = name
            break

    return PageSetup(
        paper_size=matched_paper,
        orientation=orient,
        margin_top_cm=to_cm(section.top_margin),
        margin_bottom_cm=to_cm(section.bottom_margin),
        margin_left_cm=to_cm(section.left_margin),
        margin_right_cm=to_cm(section.right_margin),
        header_distance_cm=to_cm(section.header_distance),
        footer_distance_cm=to_cm(section.footer_distance),
        center_horizontally=False,
        center_vertically=False,
    )


def read_page_setup_xlsx(xlsx_path: Path) -> PageSetup:
    """Read page setup from an Excel sheet (.xlsx)."""
    import openpyxl

    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    ws = wb.active or wb.worksheets[0]
    ws.page_setup._parent = ws

    paper_code = ws.page_setup.paperSize
    matched_paper = "Custom"
    if paper_code is not None:
        p_str = str(paper_code)
        for name, code in _OPENPYXL_PAPER.items():
            if code == p_str:
                matched_paper = name
                break

    orient = ws.page_setup.orientation or "portrait"
    if isinstance(orient, str):
        orient = orient.lower()

    margin_top = inch_to_cm(ws.page_margins.top)
    margin_bottom = inch_to_cm(ws.page_margins.bottom)
    margin_left = inch_to_cm(ws.page_margins.left)
    margin_right = inch_to_cm(ws.page_margins.right)
    header = inch_to_cm(ws.page_margins.header)
    footer = inch_to_cm(ws.page_margins.footer)

    fit_to_page = ws.page_setup.fitToPage or False
    fit_w = ws.page_setup.fitToWidth or 0
    fit_h = ws.page_setup.fitToHeight or 0

    h_center = ws.print_options.horizontalCentered or False
    v_center = ws.print_options.verticalCentered or False
    gridlines = ws.print_options.gridLines or False

    # Repeat rows
    repeat_rows = 0
    if ws.print_title_rows:
        try:
            # e.g., "1:2"
            parts = ws.print_title_rows.split(":")
            if len(parts) == 2:
                repeat_rows = int(parts[1])
        except Exception:
            pass

    wb.close()

    return PageSetup(
        paper_size=matched_paper,
        orientation=orient,
        margin_top_cm=margin_top,
        margin_bottom_cm=margin_bottom,
        margin_left_cm=margin_left,
        margin_right_cm=margin_right,
        header_distance_cm=header,
        footer_distance_cm=footer,
        center_horizontally=h_center,
        center_vertically=v_center,
        fit_to_page=fit_to_page,
        fit_to_width=fit_w,
        fit_to_height=fit_h,
        print_gridlines=gridlines,
        repeat_rows=repeat_rows,
    )


def read_page_setup_pptx(pptx_path: Path) -> PageSetup:
    """Read page setup (slide size) from PowerPoint (.pptx)."""
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    w_cm = round(prs.slide_width / 360000.0, 2)
    h_cm = round(prs.slide_height / 360000.0, 2)

    orient = "landscape" if w_cm > h_cm else "portrait"
    port_w, port_h = (w_cm, h_cm) if orient == "portrait" else (h_cm, w_cm)

    matched_paper = "Custom"
    for name, (pw, ph) in PAPER_SIZES.items():
        if abs(pw - port_w) < 0.25 and abs(ph - port_h) < 0.25:
            matched_paper = name
            break

    return PageSetup(
        paper_size=matched_paper,
        orientation=orient,
        margin_top_cm=0.0,
        margin_bottom_cm=0.0,
        margin_left_cm=0.0,
        margin_right_cm=0.0,
        header_distance_cm=0.0,
        footer_distance_cm=0.0,
        center_horizontally=False,
        center_vertically=False,
    )


def read_page_setup(input_path: str | Path) -> PageSetup:
    """Detect file type and extract original PageSetup settings."""
    path = Path(input_path)
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            return read_page_setup_docx(path)
        elif ext == ".xlsx":
            return read_page_setup_xlsx(path)
        elif ext == ".pptx":
            return read_page_setup_pptx(path)
    except Exception as exc:
        logger.warning("Error reading original page setup for %s: %s", path.name, exc)
    return PageSetup()
